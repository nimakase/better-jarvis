"""潜客链路测试 —— 从选节点到出表的回归网 + 与信号库解耦的守门。

为什么值得钉死：整条链的价值全在【排序口径】和【失败不浪费】两件事上，
而这两件事都是纯确定性逻辑，本来就该被测住：

  1. crm_state 派生 + 多级排序 + 「已认领压底」——排错了，销售拿到的名单顺序就是错的。
  2. HubSpot 未登录 → 存盘-通知-续跑（v0.5）：不出半成品名单、不推进树、
     登录后重跑直接续这批（不重花 LLM 生成的钱）。
  3. 潜客树推进：跑完一个叶子标 done，下次必须换到下一个，不能原地打转；
     全部跑完是干净收尾（StopWorkflow），不是故障。
  4. v0.5 瘦身守门：assemble / 历史线（seen_before）/ crm_state=pending 都已删除，
     不许悄悄长回来。

pipeline 的 match_fn 本来就是可注入的（就是为了测试留的口子），所以全程 mock，
不碰 playwright/浏览器/网络。openpyxl 缺失时自动跳过写表那节，其余照跑。
"""
import asyncio
import json
import sys
import tempfile
from pathlib import Path

JARVIS = str(Path(__file__).resolve().parent.parent)
sys.path.insert(0, JARVIS)

from prospecting import generation as gen    # noqa: E402
from prospecting import pipeline as pl       # noqa: E402

fails = []


def check(name, cond):
    print(("PASS " if cond else "FAIL ") + name)
    if not cond:
        fails.append(name)


TMP = Path(tempfile.mkdtemp(prefix="jarvis_prospect_test_"))


# ── 1. crm_state 派生（契约 §3）────────────────────────────────────────────
check("matched + 有 owner → owned", pl.derive_crm_state("matched", "Alice") == "owned")
check("matched + owner 空 → unowned", pl.derive_crm_state("matched", "") == "unowned")
check("matched + owner 全空格 → unowned", pl.derive_crm_state("matched", "   ") == "unowned")
check("matched_owner_empty → unowned",
      pl.derive_crm_state("matched_owner_empty", "") == "unowned")
check("no_match → new", pl.derive_crm_state("no_match", "") == "new")
check("multiple_exact_matches → review",
      pl.derive_crm_state("multiple_exact_matches", "") == "review")
check("未知/报错状态 → unknown", pl.derive_crm_state("error", "") == "unknown")
check("None 状态 → unknown", pl.derive_crm_state(None, None) == "unknown")


# ── 2. 排序键：CRM 状态 > confidence（契约 §4，v0.5 两级）────────────────────
check("CRM 顺序：全新最前", pl.CRM_ORDER["new"] < pl.CRM_ORDER["unowned"])
check("CRM 顺序：未认领 > 待核", pl.CRM_ORDER["unowned"] < pl.CRM_ORDER["review"])
check("CRM 顺序：已认领沉到最底",
      pl.CRM_ORDER["owned"] == max(pl.CRM_ORDER.values()))
check("v0.5：pending 态已移除（降级不再出半成品名单）", "pending" not in pl.CRM_ORDER)

_k = pl.sort_key
check("同 CRM 下 confidence 高的排前",
      _k({"crm_state": "new", "confidence": 90}) < _k({"crm_state": "new", "confidence": 50}))
check("CRM 状态压过 confidence（高把握的已认领仍在低把握全新之后）",
      _k({"crm_state": "new", "confidence": 10}) < _k({"crm_state": "owned", "confidence": 99}))
check("缺 confidence 不崩，按 0 处理",
      _k({"crm_state": "new"}) == (pl.CRM_ORDER["new"], 0.0))
check("未知 crm_state 落兜底位（排在已认领之前）",
      pl.CRM_ORDER["owned"] > _k({"crm_state": "wat"})[0])
check("v0.5：排序键只剩两级（seen_before 已随历史线删除）",
      len(_k({"crm_state": "new", "confidence": 50, "seen_before": True})) == 2)
check("weight 合成数已移除", not hasattr(pl, "compute_weight"))
check("意向分档已移除", not hasattr(pl, "intent_tier"))
check("意向基线常量已移除", not hasattr(pl, "INTENT_BASELINE"))


# ── 3. 富化 + 排序：已认领必须整体压底（契约 §4 核心）──────────────────────
MATCHES = {
    "SureButOwned":  {"status": "matched", "owner": "Alice"},   # 把握最高但已被认领
    "WeakNew":       {"status": "no_match", "owner": ""},       # 把握低但全新
    "SureNew":       {"status": "no_match", "owner": ""},       # 把握高且全新 → 该第一
    "InCrmUnowned":  {"status": "matched", "owner": ""},        # 在库未认领
    "Ambiguous":     {"status": "multiple_exact_matches", "owner": ""},
}


def mock_match(company, domain):
    return MATCHES.get(company, {"status": "no_match", "owner": ""})


records = [
    {"company_name": "SureButOwned", "website": "owned.com", "confidence": 99},
    {"company_name": "WeakNew", "website": "weak.com", "confidence": 40},
    {"company_name": "SureNew", "website": "sure.com", "confidence": 95},
    {"company_name": "InCrmUnowned", "website": "incrm.com", "confidence": 90},
    {"company_name": "Ambiguous", "website": "amb.com", "confidence": 80},
]
enriched = pl.enrich_records(records, mock_match)
order = [r["company_name"] for r in enriched]

check("富化不丢条目（5 进 5 出）", len(enriched) == 5)
check("已认领公司保留在名单里（标记不删）", "SureButOwned" in order)
check("已认领压到最底部（哪怕把握 99）", order[-1] == "SureButOwned")
check("全新 + 高把握浮顶", order[0] == "SureNew")
check("同为全新时按把握降序", order.index("SureNew") < order.index("WeakNew"))
check("全新整体排在在库未认领之前",
      order.index("WeakNew") < order.index("InCrmUnowned"))
check("rank 从 1 连续递增", [r["rank"] for r in enriched] == [1, 2, 3, 4, 5])
check("crm_state 已写回每条",
      {r["company_name"]: r["crm_state"] for r in enriched} ==
      {"SureButOwned": "owned", "WeakNew": "new", "SureNew": "new",
       "InCrmUnowned": "unowned", "Ambiguous": "review"})
check("原始字段不被富化丢掉",
      all("website" in r and "confidence" in r for r in enriched))
check("缺 confidence 也能富化，不崩",
      len(pl.enrich_records([{"company_name": "NoConf"}], mock_match)) == 1)

# 富化后的记录不该再带任何信号字段（解耦的实证）
_sig_fields = {"intent_score", "intent_tier", "weight", "surplus_signals", "signal_ids"}
check("富化结果不含任何信号字段（潜客已与信号解耦）",
      all(not (_sig_fields & set(r)) for r in enriched))


# ── 4. 单条匹配抛错不得中断整批（韧性）────────────────────────────────────
def exploding_match(company, domain):
    if company == "Bad":
        raise RuntimeError("模拟 HubSpot 单条查询炸了")
    return {"status": "no_match", "owner": ""}


mixed = pl.enrich_records(
    [{"company_name": "Good", "confidence": 80},
     {"company_name": "Bad", "confidence": 80},
     {"company_name": "Good2", "confidence": 80}],
    exploding_match)
check("单条匹配炸了，其余仍富化（3 进 3 出）", len(mixed) == 3)
check("炸掉那条降级成 unknown 而非丢弃",
      next(r for r in mixed if r["company_name"] == "Bad")["crm_state"] == "unknown")


# ── 5. v0.5 瘦身守门：assemble 与历史线不许长回来 ──────────────────────────
from prospecting import workflows as pw  # noqa: E402

check("assemble 已删除（只剩打没人读的 track 标记，无存在价值）",
      not hasattr(gen, "assemble"))
check("history.py 已整体删除",
      not (Path(JARVIS) / "prospecting" / "history.py").exists())
check("workflows 不再有 _degrade（降级出半成品名单的路已废）",
      not hasattr(pw, "_degrade"))
check("seen_before 不在 xlsx 列里", "seen_before" not in pl.COLUMNS)
check("track 不在 xlsx 列里", "track" not in pl.COLUMNS)


# ── 6. 潜客树：节点 =（类目 × 区域），同一类目连跑完三区域再换 ──────────────
# v2.0 的核心：一次只发一个区域（一次输出兼顾多区域必然厚此薄彼），且同一类目的
# 三个区域要连着跑完——这样才能拿到同类目、同口径的区域对照。中途跳走对照就散了。
TREE = TMP / "tree.json"
TREE.write_text(json.dumps({
    "regions_meta": {"EU": "欧洲（德/意/法）", "NA": "美洲", "SEA": "东南亚"},
    "sectors": [
        {"id": "s1", "label": "赛道一", "status": "pending", "children": [
            {"id": "l1", "label": "类目一", "regions": ["EU", "NA", "SEA"],
             "hs_codes": ["8471.50"], "status": "pending", "done_regions": []},
            {"id": "l2", "label": "类目二", "regions": ["EU", "NA"],
             "hs_codes": ["8517.62"], "status": "pending", "done_regions": []},
        ]},
        {"id": "s2", "label": "赛道二", "status": "pending", "children": [
            {"id": "l3", "label": "类目三", "regions": ["EU"],
             "status": "pending", "done_regions": []},
        ]},
    ],
}, ensure_ascii=False), encoding="utf-8")

n = gen.select_node(TREE)
check("首个节点是第一个类目", n["id"] == "l1")
check("一次只发【一个】区域", n["region"] == "EU" and isinstance(n["region"], str))
check("区域给的是人话名（不是 EU 代号）", n["region_label"] == "欧洲（德/意/法）")
check("带 node_key 唯一标识 (类目,区域)", n["node_key"] == "l1:EU")
check("节点带出 hs_codes 供提示词用", n["hs_codes"] == ["8471.50"])
check("节点带出赛道名", n["sector_label"] == "赛道一")
check("不再有 node_sectors（查信号用的键已随解耦移除）", "node_sectors" not in n)
check("剩余数按 (类目,区域) 组合算（3+2+1-1=5）", n["remaining_pending"] == 5)
check("预告的下一个是【同类目的下一个区域】", "类目一" in (n["next_label"] or ""))

check("mark_node_done 按区域标记", gen.mark_node_done(TREE, "l1", "EU") is True)
n2 = gen.select_node(TREE)
check("同一类目继续跑下一个区域（不跳走）", n2["id"] == "l1" and n2["region"] == "NA")
check("已跑区域被记录", n2["done_regions"] == ["EU"])
_t = json.loads(TREE.read_text(encoding="utf-8"))
check("只跑完一个区域时类目仍是 pending",
      _t["sectors"][0]["children"][0]["status"] == "pending")
check("跑过一个区域后赛道转 partial", _t["sectors"][0]["status"] == "partial")

gen.mark_node_done(TREE, "l1", "NA")
n3 = gen.select_node(TREE)
check("第三次仍是同一类目的最后一个区域", n3["id"] == "l1" and n3["region"] == "SEA")
check("最后一个区域时预告换类目", "类目二" in (n3["next_label"] or ""))

gen.mark_node_done(TREE, "l1", "SEA")
_t = json.loads(TREE.read_text(encoding="utf-8"))
check("三区域全跑完，类目才转 done",
      _t["sectors"][0]["children"][0]["status"] == "done")
n4 = gen.select_node(TREE)
check("类目跑完后才换下一个类目", n4["id"] == "l2" and n4["region"] == "EU")

check("重复标记同一区域不会重复累加", gen.mark_node_done(TREE, "l1", "SEA") is True)
check("done_regions 无重复",
      len(json.loads(TREE.read_text(encoding="utf-8"))
          ["sectors"][0]["children"][0]["done_regions"]) == 3)
check("标记不存在的叶子返回 False", gen.mark_node_done(TREE, "nope", "EU") is False)

gen.mark_node_done(TREE, "l2", "EU")
gen.mark_node_done(TREE, "l2", "NA")
_t = json.loads(TREE.read_text(encoding="utf-8"))
check("赛道全部叶子 done → 赛道 done", _t["sectors"][0]["status"] == "done")

gen.mark_node_done(TREE, "l3", "EU")
check("全部组合跑完时 select_node 返回 None", gen.select_node(TREE) is None)

# region=None 是兼容老调用的口子：一次标掉全部区域。正常路径不该走。
TREE2 = TMP / "tree2.json"
_t2 = json.loads(TREE.read_text(encoding="utf-8"))
for _s in _t2["sectors"]:
    for _l in _s["children"]:
        _l["done_regions"] = []
        _l["status"] = "pending"
TREE2.write_text(json.dumps(_t2, ensure_ascii=False), encoding="utf-8")
gen.mark_node_done(TREE2, "l1")     # 不传 region
check("region=None 时一次标掉该类目全部区域",
      json.loads(TREE2.read_text(encoding="utf-8"))
      ["sectors"][0]["children"][0]["status"] == "done")


# ── 7. 真实潜客树自检（仓库里那份必须可用）────────────────────────────────
REAL_TREE = Path(JARVIS) / "data" / "prospect_tree.json"
if REAL_TREE.exists():
    real = json.loads(REAL_TREE.read_text(encoding="utf-8"))
    leaves = [l for s in real["sectors"] for l in s.get("children", [])]
    ids = [l["id"] for l in leaves]
    check("真实树：叶子 id 无重复", len(ids) == len(set(ids)))
    check("真实树：每个叶子都有 regions", all(l.get("regions") for l in leaves))
    check("真实树：每个叶子都有 hs_codes", all(l.get("hs_codes") for l in leaves))
    check("真实树：有 regions_meta 供展开人话区域名", bool(real.get("regions_meta")))
    check("真实树：每个类目都有 done_regions（区域各自推进）",
          all(isinstance(l.get("done_regions"), list) for l in leaves))
    _combos = sum(len(l["regions"]) for l in leaves)
    check(f"真实树：(类目×区域) 组合数 = {_combos}", _combos == len(leaves) * 3)
    check("真实树：所有 regions 代号都在 regions_meta 里",
          all(r in real["regions_meta"] for l in leaves for r in l["regions"]))
    check("真实树：select_node 能选出节点", gen.select_node(REAL_TREE) is not None)
    check("真实树不再有 signal_sectors（与信号库解耦）",
          not any("signal_sectors" in l for l in leaves))
    check("真实树不再声明 signal_vocab", "signal_vocab" not in real)
else:
    check("真实树 data/prospect_tree.json 存在", False)


# ── 8. 工作流：存盘-通知-续跑 + 树跑完干净收尾（v0.5 降级路径重做）──────────
# 这是 v0.5 的核心变更，全流程 mock 走真 run_workflow：
#   a) 未登录：不出表、不推进树、候选存盘、发通知、run.status 仍是 ok（不算故障）
#   b) 登录后重跑：不再调 generate（省 LLM 钱），直接续存盘批 → 出表 → 推进树 → 清存盘
#   c) 树全部跑完：干净收尾 + 提示 reset，不是 RuntimeError
from core import workflow as wf  # noqa: E402

WTREE = TMP / "wtree.json"
WTREE.write_text(json.dumps({
    "regions_meta": {"EU": "欧洲"},
    "sectors": [{"id": "s1", "label": "赛道", "status": "pending", "children": [
        {"id": "w1", "label": "类目W", "regions": ["EU"], "hs_codes": ["1"],
         "status": "pending", "done_regions": []},
    ]}],
}, ensure_ascii=False), encoding="utf-8")

PENDING = TMP / "pending_batch.json"
notices: list[str] = []
outputs: list[dict] = []
gen_calls = {"n": 0}

CANDS = [{"company_name": "Alpha", "website": "alpha.com", "confidence": 90},
         {"company_name": "Beta", "website": "beta.com", "confidence": 70}]


def fake_generate(node):
    gen_calls["n"] += 1
    return [dict(c) for c in CANDS]


def fake_output(ctx):
    outputs.append({"n": len(ctx.get("enrich") or []), "node": ctx.get("node")})
    return {"path": "fake.xlsx"}


def build(preflight_ok: bool):
    return pw.build_prospect_workflow(
        tree_path=WTREE,
        generate_fn=fake_generate,
        preflight_fn=lambda: preflight_ok,
        match_fn=mock_match,
        output_fn=fake_output,
        pending_path=PENDING,
        notify_fn=notices.append,
    )


# a) 未登录
run1 = asyncio.run(wf.run_workflow("t1", build(preflight_ok=False)))
check("未登录：run.status 是 ok（干净收尾，不算故障）", run1.status == "ok")
check("未登录：停在 preflight", run1.stopped_at == "preflight")
check("未登录：候选已存盘", PENDING.exists())
check("未登录：存盘批内容完整",
      len((pw.load_pending(PENDING) or {}).get("records", [])) == 2)
check("未登录：发了通知且提到存盘和登录",
      len(notices) == 1 and "存盘" in notices[0] and "登录" in notices[0])

# 实跑翻过车的回归钉：HubSpot 是 SPA，goto 返回时壳还在渲染，单拍 detect_auth_state
# 只能拿到 unknown → 在 cookie 完全有效时误判未登录。探测必须轮询 wait_for_auth_ready。
_hs_src = (Path(JARVIS) / "prospecting" / "hubspot_session.py").read_text(encoding="utf-8")
check("hubspot_session 探测登录态用轮询 wait_for_auth_ready（单拍必误判）",
      _hs_src.count("wait_for_auth_ready") >= 2)   # acquire 与 check 两处都要
check("hubspot_session 不再单拍 detect_auth_state 判 headless 登录态",
      "if probe.detect_auth_state() == \"ok\":\n            lg.info(\"hubspot_session: 已登录" not in _hs_src)
check("未登录：没有出表", outputs == [])
check("未登录：树没有被推进（旧降级路径的 bug 不再犯）",
      json.loads(WTREE.read_text(encoding="utf-8"))
      ["sectors"][0]["children"][0]["done_regions"] == [])
check("未登录：generate 只调了一次", gen_calls["n"] == 1)

# b) 登录后重跑 → 续存盘批
run2 = asyncio.run(wf.run_workflow("t2", build(preflight_ok=True)))
check("续跑：run.status ok 且没有中途停", run2.status == "ok" and run2.stopped_at is None)
check("续跑：不再调 generate（不重花 LLM 的钱）", gen_calls["n"] == 1)
check("续跑：出了表且条数对", len(outputs) == 1 and outputs[0]["n"] == 2)
check("续跑：接续的是存盘批的节点", outputs[0]["node"]["id"] == "w1")
check("续跑：树被推进", json.loads(WTREE.read_text(encoding="utf-8"))
      ["sectors"][0]["children"][0]["done_regions"] == ["EU"])
check("续跑：存盘批已清掉", not PENDING.exists())
check("续跑：enrich 真跑了匹配（排序正确，把握高的在前）",
      [r["company_name"] for r in run2.context["enrich"]] == ["Alpha", "Beta"])

# c) 树全部跑完 → 干净收尾
notices.clear()
run3 = asyncio.run(wf.run_workflow("t3", build(preflight_ok=True)))
check("树扫完：run.status ok（正常结束不是故障）", run3.status == "ok")
check("树扫完：停在 select", run3.stopped_at == "select")
check("树扫完：通知里提示了 reset 复位", len(notices) == 1 and "reset" in notices[0])
check("树扫完：没有再出表", len(outputs) == 1)

# 正常整跑（无存盘批、已登录）：generate → 出表 → 推进
for _l in (_wt := json.loads(WTREE.read_text(encoding="utf-8")))["sectors"][0]["children"]:
    _l["done_regions"] = []
    _l["status"] = "pending"
WTREE.write_text(json.dumps(_wt, ensure_ascii=False), encoding="utf-8")
run4 = asyncio.run(wf.run_workflow("t4", build(preflight_ok=True)))
check("正常整跑：全程无停顿", run4.status == "ok" and run4.stopped_at is None)
check("正常整跑：generate 被调用", gen_calls["n"] == 2)
check("正常整跑：checkpoint 落了盘又被清掉", not PENDING.exists())

# 生产形态：preflight/enrich 是 async 包装（浏览器活在专属线程，事件循环只 await）。
# 这是真实调用路径——sync Playwright 在事件循环线程里直接拒绝运行，
# make_hubspot_runtime 因此返回 async 的 preflight/enrich。步骤必须接得住 awaitable。
for _l in (_wt2 := json.loads(WTREE.read_text(encoding="utf-8")))["sectors"][0]["children"]:
    _l["done_regions"] = []
    _l["status"] = "pending"
WTREE.write_text(json.dumps(_wt2, ensure_ascii=False), encoding="utf-8")
async_calls = {"pre": 0, "enr": 0}


async def async_preflight():
    async_calls["pre"] += 1
    return True


async def async_enrich(records):
    async_calls["enr"] += 1
    return pl.enrich_records(records, mock_match)


run5 = asyncio.run(wf.run_workflow("t5", pw.build_prospect_workflow(
    tree_path=WTREE, generate_fn=fake_generate, preflight_fn=async_preflight,
    enrich_fn=async_enrich, output_fn=fake_output, pending_path=PENDING,
    notify_fn=notices.append)))
check("async preflight/enrich（生产形态）整跑成功",
      run5.status == "ok" and run5.stopped_at is None)
check("async preflight 真被 await 了", async_calls["pre"] == 1)
check("async enrich 真被 await 了且结果进 ctx",
      async_calls["enr"] == 1 and len(run5.context["enrich"]) == 2)

# match_fn / enrich_fn 至少要给一个——都缺是接线错误，必须当场炸而不是跑到一半
try:
    pw.build_prospect_workflow(tree_path=WTREE, generate_fn=fake_generate,
                               preflight_fn=lambda: True, output_fn=fake_output,
                               pending_path=PENDING)
    _raised = False
except ValueError:
    _raised = True
check("match_fn 与 enrich_fn 都缺时明确报错", _raised)

# 步骤形状：瘦身后只剩 8 步，没有 assemble / history_mark / history_record
step_names = [s.name for s in build(True)]
check("步骤清单符合 v0.5 瘦身",
      step_names == ["select", "generate", "checkpoint", "preflight",
                     "enrich", "output", "mark_done", "clear_pending"])
check("mark_done 在 output 之后（先交付再推进）",
      step_names.index("mark_done") > step_names.index("output"))
check("checkpoint 在 preflight 之前（未登录也不丢生成结果）",
      step_names.index("checkpoint") < step_names.index("preflight"))

# 存盘批的健壮性
check("load_pending 对不存在的文件返回 None", pw.load_pending(TMP / "nope.json") is None)
_bad = TMP / "bad.json"
_bad.write_text("{{{", encoding="utf-8")
check("load_pending 对坏文件返回 None 而不是崩", pw.load_pending(_bad) is None)

# 产出文件必须能被 run_workflow 工具捞到并自动发到对话——
# 旧实现返回里连路径都没有，模型想发也不知道发什么，文件从结构上就到不了人手里。
from connectors.workflow_tools import extract_output_file  # noqa: E402

_ff = TMP / "prospects_demo.xlsx"
_ff.write_text("x", encoding="utf-8")
check("extract_output_file 捞到 output.path",
      extract_output_file({"output": {"path": str(_ff)}}) == str(_ff))
check("path=None（未出表）返回 None", extract_output_file({"output": {"path": None}}) is None)
check("文件不存在返回 None",
      extract_output_file({"output": {"path": str(TMP / 'ghost.xlsx')}}) is None)
check("空 context 不崩", extract_output_file({}) is None)

# 钉死：screener 自带 file_download 动作，不再在文本里乞求模型转发
_scr = (Path(JARVIS) / "skills" / "oem_ems_screener" / "tool.py").read_text(encoding="utf-8")
check("screener 直接挂 file_download 动作发文件", "file_download(str(path)" in _scr)
check("screener 不再乞求模型调 send_file_to_chat",
      "请调用 send_file_to_chat" not in _scr)


# ── 9. 潜客与信号彻底解耦（契约 v0.4 §10）─────────────────────────────────
# 源码层面钉死：整个 prospecting 包不许 import intel 下的任何模块。
# 用 AST 而不是文本匹配：docstring 里正当地提到了 "不再 import signal_library"，
# 按字符串找会把说明文字本身判成违规（第一版就是这么误报的）。
import ast as _ast  # noqa: E402

_offenders = []
for _f in sorted((Path(JARVIS) / "prospecting").glob("*.py")):
    _tree = _ast.parse(_f.read_text(encoding="utf-8"), filename=str(_f))
    for _n in _ast.walk(_tree):
        if isinstance(_n, _ast.Import):
            for _a in _n.names:
                if _a.name.split(".")[0] == "intel":
                    _offenders.append(f"{_f.name}:{_n.lineno}")
        elif isinstance(_n, _ast.ImportFrom):
            if (_n.module or "").split(".")[0] == "intel":
                _offenders.append(f"{_f.name}:{_n.lineno}")
check(f"prospecting 包不 import intel.*（越界：{_offenders or '无'}）",
      not _offenders)

check("意向打分函数已移除", not hasattr(gen, "attach_intent"))
check("信号类型标签表已移除", not hasattr(gen, "TYPE_LABEL"))
check("机会轨构造函数早已移除", not hasattr(gen, "build_opportunity_track"))


# ── 10. 点名公司走日报 + days 窗口（信号侧不受瘦身影响）───────────────────
SDB = TMP / "signals.db"
from intel import signal_library as _sl  # noqa: E402
_sl.init_db(SDB)
_sl.ingest_signals([
    {"scope": "sector", "signal_type": "oversupply", "sectors": ["industrial"],
     "summary": "工业自动化订单回吐", "severity": 4, "surplus_implication": 3,
     "confidence": 4, "source_url": "https://ex.com/s1"},
    {"scope": "company", "signal_type": "closure", "sectors": ["industrial"],
     "summary": "Doomed Robotics 关闭德国工厂并清库", "severity": 5,
     "surplus_implication": 5, "confidence": 4, "source_url": "https://ex.com/s2",
     "companies": [{"company_name": "Doomed Robotics", "website": "doomed.example",
                    "country": "Germany", "note": "关厂清理 MCU 库存"}]},
], db_path=SDB)

pointed = _sl.company_pointed_signals(db_path=SDB, days=14)
check("日报能取到点名公司", len(pointed) == 1)
check("点名公司带出公司名", pointed[0]["company_name"] == "Doomed Robotics")
check("点名公司带出 summary 供日报渲染", bool(pointed[0].get("summary")))

# 窗口是这个查询唯一的收敛手段（它没有强度衰减）——窗口失效 = 名单/日报被历史累积淹没
old = _sl.company_pointed_signals(db_path=SDB, days=14, as_of="2026-12-31")
check("超出 days 窗口的老信号被排除", old == [])
check("days=None 时不限窗口（向后兼容）",
      len(_sl.company_pointed_signals(db_path=SDB, days=None)) == 1)

from intel import report as _rp  # noqa: E402
from intel import workflow_defs as _wd  # noqa: E402
html = _rp.build_report_html(db_path=SDB, days=14)
check("日报出现「点名公司」板块", "点名公司" in html)
check("日报正文含点名到的公司名", "Doomed Robotics" in html)
check("日报含该公司的余料理由", "关厂清理" in html)
check("概览卡片数与 CSS 列数一致（5 列）",
      html.count('class="card-label"') == 5 and "repeat(5,1fr)" in html)

EMPTY = TMP / "empty.db"
_sl.init_db(EMPTY)
empty_html = _rp.build_report_html(db_path=EMPTY, days=14)
check("无点名公司时板块自动隐藏", "点名公司（金线索）" not in empty_html)


# ── 11. 生成截断不得整批归零（曾经的真 bug）──────────────────────────────
from core.json_salvage import salvage_json_array, looks_truncated  # noqa: E402

TRUNCATED = ('好的，以下是候选公司：\n[\n'
             '  {"company_name":"Alpha GmbH","website":"alpha.de","country":"Germany",'
             '"components":["MCU","SiC_IGBT"],"contact_rationale":"机械臂主控换代，MCU 余料"},\n'
             '  {"company_name":"Beta Ltd","website":"beta.co.uk","country":"UK",'
             '"components":["MCU"],"contact_rationale":"产线控制器 EOL"},\n'
             '  {"company_name":"Gamma BV","website":"gamma.nl","country":"Nether')

salvaged = salvage_json_array(TRUNCATED)
check("截断输出仍救回已完整的候选（不再整批归零）", len(salvaged) == 2)
check("救回的内容正确", [r["company_name"] for r in salvaged] == ["Alpha GmbH", "Beta Ltd"])
check("被截断的最后一条被丢弃",
      not any(r.get("company_name") == "Gamma BV" for r in salvaged))
check("looks_truncated 认出截断", looks_truncated(TRUNCATED) is True)

# 旧的朴素解析在同一输入下是什么下场——钉住这个对比，防止有人改回去
naive = []
_i, _j = TRUNCATED.find("["), TRUNCATED.rfind("]")
if _i >= 0 and _j > _i:
    try:
        naive = json.loads(TRUNCATED[_i:_j + 1])
    except Exception:
        naive = []
check("对照：朴素解析在同样输入下确实归零", naive == [])

COMPLETE = '[{"company_name":"Solo","website":"solo.com"}]'
check("正常完整输出照常解析", len(salvage_json_array(COMPLETE)) == 1)
check("looks_truncated 不误报完整输出", looks_truncated(COMPLETE) is False)
check("rationale 里含花括号/引号也不会算错深度",
      len(salvage_json_array(
          r'[{"company_name":"X","contact_rationale":"用了 {MCU} 和 \"IGBT\""}]')) == 1)
check("完全不是 JSON 时返回空而不是崩", salvage_json_array("模型今天罢工了") == [])
check("空输入返回空", salvage_json_array("") == [])
check("采集轨复用同一实现（逻辑不再分叉两份）",
      _wd._parse_signal_array is salvage_json_array)

# 生成预算必须显著高于聊天护栏，否则 100 家候选必被切
import config as _cfg  # noqa: E402
check(f"生成预算({pw._GEN_MAX_TOKENS}) 显著高于聊天护栏({_cfg.MAX_TOKENS_RESPONSE})",
      pw._GEN_MAX_TOKENS >= 3 * _cfg.MAX_TOKENS_RESPONSE)
check("生成预算与采集预算同档（两条大批量生成的路口径一致）",
      pw._GEN_MAX_TOKENS == _wd._COLLECT_MAX_TOKENS)
check("生成不再走 4096 的聊天护栏",
      pw._GEN_MAX_TOKENS != _cfg.MAX_TOKENS_RESPONSE)

# 提示词占位符必须真被填上——漏填会让模型收到字面量 {{HS_CODES}}
rendered = pw.render_generation_prompt(
    "赛道={{SECTOR_LABEL}} 类目={{NODE_LABEL}} HS={{HS_CODES}} 区域={{REGION}}",
    {"label": "工业机器人", "sector_label": "工业自动化", "hs_codes": ["8479.50"],
     "region": "EU", "region_label": "欧洲（德/意/法）"})
check("提示词占位符全部被替换", "{{" not in rendered)
check("区域填的是人话名而不是 EU 代号", "欧洲" in rendered and "=EU" not in rendered)
check("HS 码进了提示词", "8479.50" in rendered)
check("缺 hs_codes 时给出占位而不是空白",
      "（未指定）" in pw.render_generation_prompt("HS={{HS_CODES}}", {"label": "x"}))


# ── 12. 汇总条目不得冒充公司 + 客户画像口径（契约 v0.3）────────────────────
parsed = [
    {"company_name": "Alpha", "website": "a.com", "category": "oem", "confidence": 90},
    {"company_name": "Beta", "website": "b.com", "category": "ems", "confidence": 70},
    {"_meta": True, "coverage_note": "欧洲已基本覆盖", "exhausted": True},
]
recs, meta = pw.split_meta(parsed)
check("_meta 条目被剥离，不进名单", len(recs) == 2)
check("剥离后名单里没有无名条目", all(r.get("company_name") for r in recs))
check("coverage_note 被接住", meta.get("coverage_note") == "欧洲已基本覆盖")
check("exhausted 被接住", meta.get("exhausted") is True)

noname = [{"company_name": "OK", "website": "ok.com"},
          {"note": "以上就是全部了"},           # 模型偶尔塞的说明性对象
          {"company_name": "   ", "website": "x.com"}]  # 空白名字
recs2, _ = pw.split_meta(noname)
check("没有 company_name 的字典也被挡掉", len(recs2) == 1)
check("纯空白名字同样被挡掉", recs2[0]["company_name"] == "OK")
check("非字典元素被忽略而不是崩", pw.split_meta(["乱码", 42, None])[0] == [])
check("空输入不崩", pw.split_meta([]) == ([], {}))

# 新增三列必须真的在表上，否则"降 confidence 而不丢弃"的策略等于没有——
# 低把握的条目会和高把握的长得一模一样
for col in ("category", "confidence", "evidence"):
    check(f"xlsx 有 {col} 列（契约 v0.3）", col in pl.COLUMNS)
check("contact_rationale 仍在（未被误删）", "contact_rationale" in pl.COLUMNS)

# 提示词口径：EMS/模组厂曾被「整机 OEM」措辞误伤，这是真实丢客户的 bug
PROMPT_GEN = Path(JARVIS) / "intel" / "prospect_generation_prompt.md"
ptxt = PROMPT_GEN.read_text(encoding="utf-8")
check("提示词把 ems 列为合格类别", "`ems`" in ptxt)
check("提示词把 module 列为合格类别", "`module`" in ptxt)
check("提示词明确提到 SMT/PCBA 代工", "PCBA" in ptxt)
check("提示词要求给 evidence", "evidence" in ptxt)
check("提示词要求 confidence 校准", "confidence" in ptxt)
check("提示词不再让模型去猜『为什么现在有余料』",
      "不要写\"为什么现在可能有余料\"" in ptxt or "不要写“为什么现在可能有余料”" in ptxt)
check("提示词明确『查不到不等于不合格』",
      "absence of evidence" in ptxt)
check("提示词给了 _meta 汇总条目的合法位置", "_meta" in ptxt)
check("规模改为分档、不再卡人数", "size_tier" in ptxt and "不要去查员工人数当门槛" in ptxt)
check("五档都在提示词里",
      all(f"`{t}`" in ptxt for t in ("startup", "small", "mid", "large", "giant")))
check("明确排除巨头的理由是『够不着』而非人多", "够不着" in ptxt)
check("给了边界情况的判断依据（3000 人但采购分散仍要）", "3000 人" in ptxt)
# v0.5：中国大陆与香港硬排除（开发难度过高），且判断标准是母公司/总部所在地
check("提示词硬排除中国大陆与香港公司",
      "中国大陆与香港的公司一律排除" in ptxt)
check("排除口径按母公司/总部判，不按注册地", "总部所在地" in ptxt)
# size_tier 是给模型的筛选口径，不是给人的数据——不该出现在报表列里
check("size_tier 不进 xlsx 列", "size_tier" not in pl.COLUMNS)
check("size_basis 不进 xlsx 列", "size_basis" not in pl.COLUMNS)
check("提示词占位符与渲染函数对齐",
      all(k in ptxt for k in ("{{NODE_LABEL}}", "{{SECTOR_LABEL}}",
                              "{{HS_CODES}}", "{{REGION}}")))
rendered_real = pw.render_generation_prompt(ptxt, {
    "label": "工业机器人与机械臂", "sector_label": "工业自动化与工厂设备",
    "hs_codes": ["8479.50"], "region": "EU", "region_label": "欧洲"})
check("真实提示词渲染后无残留占位符", "{{" not in rendered_real)


# ── 13. 信号过期必须真的有人调（status 列不能是假的）──────────────────────
EDB = TMP / "expire.db"
_sl.init_db(EDB)
# closure 半衰期 120 天，3× = 360 天才过期；oversupply 45 天，3× = 135 天
_sl.ingest_signals([
    {"scope": "company", "signal_type": "closure", "sectors": ["industrial"],
     "summary": "陈年旧闻：某厂两年前关闭", "severity": 5, "surplus_implication": 5,
     "confidence": 4, "source_url": "https://ex.com/old",
     "companies": [{"company_name": "Ancient Co", "website": "ancient.example"}]},
], db_path=EDB, collected_date="2024-01-01")
_sl.ingest_signals([
    {"scope": "company", "signal_type": "closure", "sectors": ["industrial"],
     "summary": "新鲜事：某厂本周关闭", "severity": 5, "surplus_implication": 5,
     "confidence": 4, "source_url": "https://ex.com/new",
     "companies": [{"company_name": "Fresh Co", "website": "fresh.example"}]},
], db_path=EDB, collected_date="2026-07-17")

before = _sl.company_pointed_signals(db_path=EDB, as_of="2026-07-18")
check("过期前：新旧信号都还在（status 全是 active）", len(before) == 2)

n_expired = _sl.expire_stale(db_path=EDB, as_of="2026-07-18")
check("expire_stale 把陈年信号标成 expired", n_expired == 1)

after = _sl.company_pointed_signals(db_path=EDB, as_of="2026-07-18")
check("过期后只剩新鲜信号", len(after) == 1)
check("留下的是新的那条", after[0]["company_name"] == "Fresh Co")
check("重复调用不会重复计数（幂等）",
      _sl.expire_stale(db_path=EDB, as_of="2026-07-18") == 0)

# 再次采到同一条老信号 → 合并会刷新 date_collected 并复活成 active
_sl.ingest_signals([
    {"scope": "company", "signal_type": "closure", "sectors": ["industrial"],
     "summary": "陈年旧闻：某厂两年前关闭", "severity": 5, "surplus_implication": 5,
     "confidence": 4, "source_url": "https://ex.com/old",
     "companies": [{"company_name": "Ancient Co", "website": "ancient.example"}]},
], db_path=EDB, collected_date="2026-07-18")
revived = _sl.company_pointed_signals(db_path=EDB, as_of="2026-07-18")
check("再次采到会复活（反复出现的行情不会被误清）", len(revived) == 2)

# 光有函数不够，必须真的挂在工作流里——否则又变回死代码
steps = _wd._build_signal_collection()
names = [s.name for s in steps]
check("采集工作流包含 expire 步", "expire" in names)
check("expire 排在 ingest 之后（本次采的不会被误伤）",
      names.index("expire") > names.index("ingest"))
check("expire 失败不拖垮采集（on_error=skip）",
      next(s for s in steps if s.name == "expire").on_error == "skip")


# ── 14. 写 xlsx（openpyxl 缺失则跳过）──────────────────────────────────────
try:
    import openpyxl  # noqa: F401
    has_openpyxl = True
except ImportError:
    has_openpyxl = False

if has_openpyxl:
    check("列表渲染成分号串", pl._as_text(["涨价", "EOL"]) == "涨价; EOL")
    check("None 渲染成空", pl._as_text(None) == "")

    out = TMP / "out.xlsx"
    pl.write_xlsx(enriched, str(out))
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    check("xlsx 落盘", out.exists())
    check("表头即契约 §2 列", [c.value for c in ws[1]] == pl.COLUMNS)
    check("数据行数对得上", ws.max_row == 1 + len(enriched))
    check("冻结在表头下方", ws.freeze_panes == "A2")
    import inspect
    check("v0.5：write_xlsx 不再有 banner 参数（随降级路径一起删除）",
          "banner" not in inspect.signature(pl.write_xlsx).parameters)
else:
    print("SKIP xlsx 相关检查（未安装 openpyxl）")


# ── 汇总 ─────────────────────────────────────────────────────────────────
import shutil  # noqa: E402
shutil.rmtree(TMP, ignore_errors=True)

print("\n" + "=" * 52)
if fails:
    print(f"❌ {len(fails)} 项失败：")
    for f in fails:
        print("   - " + f)
    sys.exit(1)
print("✅ 潜客链路全部检查通过")
