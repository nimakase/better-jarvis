"""
潜客流水线 · 富化与排序层（Prospect pipeline enrichment + ranking）

把潜客候选逐条过 HubSpot 匹配 → 派生 crm_state → 多级排序（已认领置底）→ 输出 xlsx。

契约见 intel/prospect_pipeline_contract.md。

**v0.4：与信号库完全解耦。** 本模块（及整个 prospecting 包）不再读信号库。
潜客名单服务的是 **0→1 大范围开发**——每天拿一批新公司去建立关系，周期以月计；
而信号的时效是几周，且是赛道级/元件级的，一整个赛道的公司拿到同一个分数，
在一批候选内部几乎没有区分度，只会让不同日子的名单互相不可比。两者节奏与粒度都不匹配。
信号的真实用途在别处：判断手里的货要不要压，以及回头找买过某元件的老客户——
都是**存量**判断，走情报侧（市场情报日报），不在这条线上。

**v0.5 瘦身：**
  - 移除 seen_before 排序级与列——历史库整条移除（重复率低，不值一个存储层）。
  - 移除 crm_state=pending——降级路径改成「存盘待续跑」，不再出未匹配的名单
    （见 workflows.py），于是"半成品状态"不复存在。
  - write_xlsx 去掉 banner 参数——它唯一的用途是降级提示，随降级路径一起消失。

匹配函数可注入（match_fn）：
  - 生产：make_hubspot_match_fn(browser) 包住爬虫内核 process_one。
  - 测试：传一个 mock，返回 {"status":..., "owner":...}，无需浏览器。

本模块只用标准库；写 xlsx 时才惰性 import openpyxl。
"""
from __future__ import annotations

from typing import Callable, Optional

# CRM 状态的排序优先级（contract §4）——数字小的排前面。
# 这是 0→1 开发的自然优先级：没人碰过的最有价值，已被同事认领的压底但不删。
CRM_ORDER = {
    "new": 0,        # HubSpot 里没有 → 全新潜客，最该打
    "unowned": 1,    # 在 CRM 但没人认领 → 值得跟
    "review": 2,     # 多重命中，待人工核
    "unknown": 4,    # 匹配未能确定
    "owned": 9,      # 已被销售认领 —— 保留但压到最底
}

MatchFn = Callable[[str, str], dict]


# ────────────────────────── 派生与计算 ──────────────────────────

def derive_crm_state(status: Optional[str], owner: Optional[str]) -> str:
    """把 matcher 返回的 status/owner 映射成 crm_state（contract §3）。"""
    owner = (owner or "").strip()
    if status == "matched":
        return "owned" if owner else "unowned"
    if status == "matched_owner_empty":
        return "unowned"
    if status == "no_match":
        return "new"
    if status == "multiple_exact_matches":
        return "review"
    return "unknown"


def sort_key(rec: dict) -> tuple:
    """名单排序键（contract §4，v0.5）。两级，每一级都是你看得懂的量。

        1. crm_state  —— 全新 > 在库未认领 > 待核 > …… > 已认领（压底）
        2. confidence —— 同等状态下，把握大的先打（降序）

    为什么不再合成一个 weight 数：老版本是 `intent_score × crm系数`，那时 intent
    是信号驱动的真实强度，乘出来有意义。现在信号已解耦，若拿 confidence 去乘
    CRM 系数凑一个小数，是**假精度**——两个不同性质的量相乘，小数点后三位没有
    任何含义，还让人误以为它精确。多级排序把每一级摊开，谁排前面、为什么排前面，
    一眼能看明白，也方便你在表里自己重排。

    （v0.5 移除了第三级 seen_before——历史库整条删除，见模块 docstring。）
    """
    return (
        CRM_ORDER.get(rec.get("crm_state"), 5),
        -float(rec.get("confidence") or 0),
    )


# ────────────────────────── 富化主流程 ──────────────────────────

# 最终记录的列顺序（contract §2）
#
# v0.3 加了 category / confidence / evidence（生成层产出）：
#   category   —— oem / ems / module，与 oem_ems_screener 同一套词表
#   confidence —— 生成层对"这家确实符合画像"的把握（0–100，诚实校准）
#   evidence   —— 依据来源；没有它就无法核实名单，只能全信或全不信
# 加这三列是因为提示词改成了"查不到就降 confidence 而不是丢弃"——
# 若不把 confidence/evidence 摆到表上，低把握的条目会和高把握的长得一模一样。
#
# v0.4 移除 weight / intent_tier / intent_score / surplus_signals / sources 五列
# （随信号解耦）；v0.5 移除 seen_before（随历史库删除）。
# v0.6 加 parent_note（生成层产出）：官网/资料明确提到的母公司/集团归属，纯软信号，
#   不做自动排除（won 客户里有独立成交的集团子公司反例，见 prospect_generation_prompt.md
#   「对分支机构不做硬排除」的讨论）——留空表示查不到归属或就是独立公司。
COLUMNS = [
    "rank", "crm_state", "hubspot_status", "hubspot_owner",
    "category", "confidence",
    "company_name", "website", "country", "components",
    "contact_rationale", "evidence", "parent_note",
]


def _as_text(value) -> str:
    if value is None:
        return ""
    # bool 要在 int 之前判：Python 里 bool 是 int 的子类。
    # 表里 True→「是」、False→空白（而不是刺眼的一列「否」）。
    if isinstance(value, bool):
        return "是" if value else ""
    if isinstance(value, (list, tuple)):
        return "; ".join(str(v) for v in value)
    return str(value)


def enrich_records(records: list[dict], match_fn: MatchFn) -> list[dict]:
    """逐条过 HubSpot 匹配 + 派生 crm_state + 排序。返回排好序、带 rank 的列表。

    输入 record 至少含 company_name；website/domain 用于匹配；
    可含 category / confidence / evidence / components 等（来自生成层）。
    单条匹配失败不中断整批：那一条标 status=error → crm_state=unknown。
    """
    enriched: list[dict] = []
    for r in records:
        company = (r.get("company_name") or "").strip()
        domain = (r.get("website") or r.get("domain") or "").strip()

        try:
            m = match_fn(company, domain) or {}
        except Exception as exc:  # 单条匹配失败不应中断整批
            m = {"status": "error", "owner": "", "error": str(exc)}

        status = m.get("status")
        owner = m.get("owner") or ""

        out = dict(r)
        out.update({
            "hubspot_status": status,
            "hubspot_owner": owner,
            "crm_state": derive_crm_state(status, owner),
        })
        enriched.append(out)

    # 多级排序（contract §4）：CRM 状态 > confidence。
    # 已认领的靠 CRM_ORDER 里的 9 自然沉底——保留但压底，不删。
    enriched.sort(key=sort_key)
    for i, rec in enumerate(enriched, 1):
        rec["rank"] = i
    return enriched


# ────────────────────────── 输出富 xlsx ──────────────────────────

def write_xlsx(records: list[dict], path: str) -> str:
    """把富记录写成 xlsx（列见 COLUMNS）。已认领行浅灰底以示压底。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "prospects"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A3A5C")
    owned_fill = PatternFill("solid", fgColor="EEEEEE")

    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")

    for rec in records:
        ws.append([_as_text(rec.get(col)) for col in COLUMNS])
        if rec.get("crm_state") == "owned":
            for c in ws[ws.max_row]:
                c.fill = owned_fill

    widths = {"company_name": 28, "website": 22, "contact_rationale": 50,
              "components": 18, "country": 14,
              "evidence": 36, "category": 10, "confidence": 10,
              "parent_note": 22}
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    wb.save(path)
    return path


# ────────────────────────── 生产：包住爬虫内核 ──────────────────────────

def make_hubspot_match_fn(browser, logger=None) -> MatchFn:
    """生产用：把已就绪（已登录、列映射已建）的 HubSpotBrowser 包成 match_fn。

    惰性 import hubspot_worker（依赖 playwright/pandas），测试路径不触发。
    """
    from prospecting import hubspot_worker as w

    def match(company: str, domain: str) -> dict:
        outcome = w.process_one(browser, company, domain)
        return {
            "status": outcome.status,
            "owner": outcome.owner,
            "company_search_result_count": outcome.company_search_result_count,
            "domain_search_result_count": outcome.domain_search_result_count,
        }

    return match
