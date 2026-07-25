"""
文档读取工具

支持格式：
  PDF    → pdfplumber 提取文字（含表格）
  DOCX   → python-docx 提取段落和表格
  XLSX   → openpyxl 提取所有 sheet
  PPTX   → python-pptx 提取每页文字
  TXT/MD → 直接读取
  CSV    → 转为 markdown 表格
  图片    → base64 发给视觉模型，返回描述/文字

使用方式（对贾维斯说）：
  "帮我读一下 C:/Users/Ned/Desktop/合同.pdf，总结主要条款"
  "读一下这个 Excel：C:/Users/Ned/Downloads/报表.xlsx"
"""

import base64
import csv
import os
from pathlib import Path
from typing import Optional

import config
from functools import partial
from core.registry import tool as _tool

# read_document 归入 fileio 组（与元工具 send_file_to_chat 同组：文件读入/发出）。
# 注意：与文档保险箱 doc_vault（group=documents）刻意分开——一个只读取，一个加密归档。
tool = partial(_tool, group="fileio")

# token 安全上限：超过此长度截断并提示
MAX_CHARS = 120_000  # ~8万 token，留足 context 给其他内容


def _truncate(text: str, path: str) -> str:
    if len(text) <= MAX_CHARS:
        return text
    return (
        text[:MAX_CHARS]
        + f"\n\n[文档过长，已截断。原文件：{path}，"
        f"已读取约前 {MAX_CHARS} 字符，共约 {len(text)} 字符。"
        f"如需后续内容请告知页码或章节。]"
    )


# ── 各格式提取函数 ────────────────────────────────────────────────────────────

def _meaningful_len(s: str) -> int:
    """统计有效字符数：只数字母/数字/CJK，忽略空白和标点。用于识别"空提取"。"""
    return sum(1 for ch in s if ch.isalnum())


def _read_pdf(path: str, pages: Optional[str] = None) -> tuple[str, int, int]:
    """本地读取 PDF（pdfplumber）。返回 (格式化文本, 页数, 有效字符数)。
    有效字符数用于判断是否"空提取"（扫描件/图片型 PDF 抽不出文字层）。

    pages: None=全部, "1"=第1页, "1-5"=第1到5页, "1,3,5"=指定页
    """
    import pdfplumber

    def parse_pages(spec: str, total: int) -> list[int]:
        result = set()
        for part in spec.split(","):
            part = part.strip()
            if "-" in part:
                a, b = part.split("-")
                result.update(range(int(a) - 1, min(int(b), total)))
            else:
                result.add(int(part) - 1)
        return sorted(result)

    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        if pages:
            indices = parse_pages(pages, total)
        else:
            indices = list(range(total))

        parts = []
        meaningful = 0
        for i in indices:
            page = pdf.pages[i]
            text = page.extract_text() or ""
            # 提取表格
            tables = page.extract_tables()
            table_text = ""
            for table in tables:
                rows = []
                for row in table:
                    rows.append(" | ".join(str(c or "").strip() for c in row))
                table_text += "\n" + "\n".join(rows)

            meaningful += _meaningful_len(text) + _meaningful_len(table_text)
            parts.append(f"--- 第 {i+1} 页 ---\n{text}{table_text}")

    formatted = f"[PDF 文档，共 {total} 页，已读 {len(indices)} 页]\n\n" + "\n\n".join(parts)
    return formatted, len(indices), meaningful


def _read_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            style = para.style.name
            prefix = ""
            if "Heading 1" in style:
                prefix = "# "
            elif "Heading 2" in style:
                prefix = "## "
            elif "Heading 3" in style:
                prefix = "### "
            parts.append(prefix + para.text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        parts.append("\n[表格]\n" + "\n".join(rows))

    return "[Word 文档]\n\n" + "\n\n".join(parts)


def _read_xlsx(path: str, sheet: Optional[str] = None) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    sheets = [sheet] if sheet and sheet in wb.sheetnames else wb.sheetnames
    parts = []

    for name in sheets:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(" | ".join(str(c) if c is not None else "" for c in row))
        parts.append(f"[Sheet: {name}]\n" + "\n".join(rows))

    return f"[Excel 文档，共 {len(wb.sheetnames)} 个 Sheet]\n\n" + "\n\n".join(parts)


def _read_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts = []

    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            parts.append(f"--- 第 {i} 页 ---\n" + "\n".join(texts))

    return f"[PowerPoint，共 {len(prs.slides)} 页]\n\n" + "\n\n".join(parts)


def _read_csv(path: str) -> str:
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        reader = csv.reader(f)
        rows = [" | ".join(row) for row in reader]
    return "[CSV 文件]\n\n" + "\n".join(rows)


def _read_text(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


async def _read_image(path: str) -> str:
    """把图片发给视觉模型，返回文字描述/OCR 结果。"""
    from openai import AsyncOpenAI

    ext = Path(path).suffix.lower().lstrip(".")
    mime_map = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
                "gif": "image/gif", "webp": "image/webp", "bmp": "image/bmp"}
    mime = mime_map.get(ext, "image/png")

    with open(path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode()

    from core.llm import get_client
    client = get_client()   # 有界超时（core/llm 单一构建点）
    resp = await client.chat.completions.create(
        model=config.CLAUDE_MODEL,
        max_tokens=2048,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                {"type": "text", "text": "请提取并描述图片中的所有文字和内容，尽量完整。"}
            ]
        }]
    )
    return "[图片内容]\n\n" + resp.choices[0].message.content


# ── 空提取判定 & 云端 OCR 兜底 ────────────────────────────────────────────────

def _pdf_is_poor(pages: int, meaningful: int) -> bool:
    """本地提取是否"基本没抽到文字"（扫描件/图片型 PDF 的典型表现）。
    阈值随页数放宽：几乎抽不到字才算 poor，避免误伤内容本就稀少的正常 PDF。"""
    return meaningful < max(20, 8 * max(pages, 1))


def _scanned_pdf_note(path: str, meaningful: int, extra: str = "") -> str:
    """本地读扫描件失败时的诚实报错（不再静默返回空壳）。"""
    base = (
        f"⚠️ 无法从该 PDF 提取到有效文字：{path}\n"
        f"本地只抽到约 {meaningful} 个有效字符——它很可能是【扫描件 / 图片型 PDF】（没有文字层）。"
    )
    return base + (("\n" + extra) if extra else "")


async def _read_pdf_via_openrouter(path: str) -> str:
    """把 PDF 交给 OpenRouter 的 file-parser 插件做云端 OCR（默认 mistral-ocr）。
    仅在文件【非敏感】且本地抽不出文字时调用。内容会离开本机——上层已做好隔离判断。"""
    import base64
    from openai import AsyncOpenAI

    data = Path(path).read_bytes()
    data_url = "data:application/pdf;base64," + base64.b64encode(data).decode()
    from core.llm import get_client
    client = get_client()   # 有界超时（core/llm 单一构建点）
    resp = await client.chat.completions.create(
        model=config.CLAUDE_MODEL_LIGHT,
        max_tokens=8192,
        messages=[{
            "role": "user",
            "content": [
                {"type": "text", "text":
                    "请完整、逐字提取该 PDF 的全部文字内容，保留原有顺序与分段；"
                    "只输出文字本身，不要总结、不要改写、不要添加任何评论或标题。"},
                {"type": "file", "file": {"filename": Path(path).name, "file_data": data_url}},
            ],
        }],
        extra_body={"plugins": [{"id": "file-parser", "pdf": {"engine": config.PDF_CLOUD_ENGINE}}]},
    )
    out = (resp.choices[0].message.content or "").strip()
    if not out:
        return f"云端 OCR 未返回内容：{path}"
    return f"[PDF · 云端 OCR 读取（OpenRouter {config.PDF_CLOUD_ENGINE}）]\n\n{out}"


async def _handle_pdf(path: str, pages: str, cloud: str) -> str:
    """PDF 读取决策：本地优先 → 扫描件按敏感度决定是否云端兜底。
    cloud: auto（默认，自动判定）/ allow（用户已同意上云）/ deny（只本地）。"""
    text, npages, meaningful = _read_pdf(path, pages or None)

    # 本地抽到了足够文字，或云端被关/被拒 → 走本地结果
    if not _pdf_is_poor(npages, meaningful):
        return _truncate(text, path)
    if cloud == "deny" or not config.PDF_CLOUD_FALLBACK:
        return _scanned_pdf_note(
            path, meaningful,
            "已按你的要求只用本地处理，未上云。如需读取，请提供含文字层的版本。")

    # 扫描件 + 云端可用：按敏感度决策
    from core import sensitivity
    if cloud == "allow":
        bucket, reason = "non_sensitive", "用户已确认可上云"
    else:
        bucket, reason = await sensitivity.classify(path, text)

    if bucket == "non_sensitive":
        try:
            return await _read_pdf_via_openrouter(path)
        except Exception as e:
            return _scanned_pdf_note(path, meaningful, f"尝试云端 OCR 失败：{type(e).__name__}: {e}")
    if bucket == "sensitive":
        return _scanned_pdf_note(
            path, meaningful,
            f"该文件被判定为【敏感】（{reason}），为保护隐私未上传云端，仅本机处理。"
            "如确需云端 OCR，请明确告知'可以上云读取'，我再处理。")
    # uncertain → 让模型去问用户；同意后用 cloud=\"allow\" 重调，拒绝用 cloud=\"deny\"
    return (
        f"❓ 该 PDF 是扫描件/图片型，本地读不出文字（{path}）。\n"
        f"是否敏感我拿不准（{reason}）。云端 OCR（OpenRouter {config.PDF_CLOUD_ENGINE}）能读，但内容会离开本机。\n"
        "【请向用户确认】是否可以把这份文件上传到云端做 OCR？\n"
        "· 用户同意 → 再次调用 read_document，参数 cloud=\"allow\"。\n"
        "· 用户不同意 → 再次调用 read_document，参数 cloud=\"deny\"（仅本地，会明确报错）。")


# ── 主工具函数 ────────────────────────────────────────────────────────────────

@tool(
    "read_document",
    (
        "读取本地文件并提取文字内容。支持 PDF、Word(.docx)、Excel(.xlsx)、"
        "PowerPoint(.pptx)、CSV、TXT、Markdown，以及 JPG/PNG 等图片（视觉识别）。"
        "用户要读取、查看、分析、总结某个文件时使用。"
        "【只读取、不保存】——若用户是要把保单/合同等存档归类，改用 ingest_document_file。"
        "path 必须是文件的完整路径。"
        "扫描件/图片型 PDF 本地读不出文字时：非敏感文件会自动走云端 OCR，敏感文件仅本机并报错；"
        "若我提示「拿不准是否敏感、需确认是否上云」，请照提示向用户确认后用 cloud 参数重新调用。"
    ),
    {
        "type": "object",
        "properties": {
            "path":  {"type": "string", "description": "文件完整路径，如 C:/Users/Ned/Desktop/合同.pdf"},
            "pages": {"type": "string", "description": "仅 PDF 有效：指定页码范围，如 '1-5' 或 '1,3,5'，留空=全部"},
            "sheet": {"type": "string", "description": "仅 Excel 有效：指定 Sheet 名称，留空=所有 Sheet"},
            "cloud": {"type": "string", "description":
                      "PDF 云端 OCR 授权：auto=自动判定(默认)；allow=用户已同意上云；deny=只用本地。"
                      "仅当我上一轮提示需要确认是否上云时，才按用户回答传 allow 或 deny。"},
        },
        "required": ["path"],
    },
)
async def read_document(path: str, pages: str = "", sheet: str = "", cloud: str = "auto") -> str:
    """
    读取文档并返回文本内容。
    path:  文件的完整路径
    pages: 仅 PDF 有效，指定页码，如 "1-5" 或 "1,3,5"，留空=全部
    sheet: 仅 XLSX 有效，指定 Sheet 名称，留空=所有 Sheet
    cloud: 仅 PDF 有效，云端 OCR 授权：auto/allow/deny
    """
    path = path.strip().strip('"').strip("'")

    if not os.path.exists(path):
        return f"文件不存在：{path}\n请确认路径正确（Windows 路径示例：C:/Users/Ned/Desktop/文件.pdf）"

    ext = Path(path).suffix.lower()
    file_size_mb = os.path.getsize(path) / 1024 / 1024

    if file_size_mb > 50:
        return f"文件过大（{file_size_mb:.1f} MB），超过 50MB 限制，请提供更小的文件或指定具体页码范围。"

    try:
        if ext == ".pdf":
            # 本地优先 + 扫描件云端兜底（按敏感度）——决策全在 _handle_pdf 里
            return _truncate(await _handle_pdf(path, pages, (cloud or "auto").lower()), path)
        elif ext in (".docx", ".doc"):
            text = _read_docx(path)
        elif ext in (".xlsx", ".xls", ".xlsm"):
            text = _read_xlsx(path, sheet or None)
        elif ext in (".pptx", ".ppt"):
            text = _read_pptx(path)
        elif ext == ".csv":
            text = _read_csv(path)
        elif ext in (".txt", ".md", ".markdown", ".log", ".json", ".yaml", ".yml"):
            text = _read_text(path)
        elif ext in (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"):
            text = await _read_image(path)
        else:
            return (
                f"不支持的文件格式：{ext}\n"
                f"支持的格式：PDF、DOCX、XLSX、PPTX、CSV、TXT、MD、JPG、PNG 等图片"
            )

        # 空提取兜底：docx/pptx/xlsx 等若几乎没抽到有效文字（可能是纯图片文档/空文件），
        # 明确报错而不是返回一个空壳让模型误以为读到了内容。
        if _meaningful_len(text) < 8:
            return (
                f"⚠️ 该文件几乎没有可提取的文字内容：{path}\n"
                "可能是纯图片文档、空文件，或内容以图片/图形形式存在（无文字层）。"
                "如果是图片扫描件，请转成图片再让我做视觉识别，或提供含文字的版本。"
            )

        return _truncate(text, path)

    except ImportError as e:
        pkg = str(e).split("'")[1] if "'" in str(e) else str(e)
        return f"缺少依赖库 {pkg}，请运行：pip install {pkg}"
    except Exception as e:
        return f"读取文件出错：{type(e).__name__}: {e}"
